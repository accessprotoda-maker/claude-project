import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.styles import Protection
from copy import copy

def copy_sheet(src_ws, dst_wb, new_name):
    dst_ws = dst_wb.create_sheet(title=new_name)
    
    for col_letter, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[col_letter].width = dim.width
        dst_ws.column_dimensions[col_letter].hidden = dim.hidden
    
    for row_num, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[row_num].height = dim.height
        dst_ws.row_dimensions[row_num].hidden = dim.hidden
    
    for merged in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merged))
    
    for row in src_ws.iter_rows(min_row=1, max_row=src_ws.max_row, max_col=src_ws.max_column):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            new_cell = dst_ws[cell.coordinate]
            new_cell.value = cell.value
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.border = copy(cell.border)
                new_cell.fill = copy(cell.fill)
                new_cell.number_format = cell.number_format
                new_cell.protection = copy(cell.protection)
                new_cell.alignment = copy(cell.alignment)
    
    if src_ws.data_validations:
        for dv in src_ws.data_validations.dataValidation:
            new_dv = DataValidation(
                type=dv.type, formula1=dv.formula1, formula2=dv.formula2,
                allow_blank=dv.allow_blank, showErrorMessage=dv.showErrorMessage,
                showInputMessage=dv.showInputMessage, errorTitle=dv.errorTitle,
                error=dv.error, promptTitle=dv.promptTitle, prompt=dv.prompt,
                operator=dv.operator,
            )
            new_dv.sqref = dv.sqref
            dst_ws.add_data_validation(new_dv)
    
    dst_ws.sheet_properties = copy(src_ws.sheet_properties)
    if src_ws.print_area:
        dst_ws.print_area = src_ws.print_area
    dst_ws.page_setup = copy(src_ws.page_setup)
    dst_ws.page_margins = copy(src_ws.page_margins)
    
    return dst_ws


def add_input_improvements(ws, sheet_type):
    if sheet_type != "fire_alarm_detail":
        return
    
    unlocked = Protection(locked=False)
    
    # Ensure all data rows (5-27) have auto-calc formulas for O and P
    for row in range(5, 28):
        for col_letter in ['O', 'P']:
            cell = ws[f"{col_letter}{row}"]
            val = cell.value
            if not (isinstance(val, str) and val.startswith('=')):
                cell.value = f'=IF(COUNTA(C{row}:N{row})=0,"","〇")'
    
    # Add number validation for sensor count columns (C-N), rows 5-27
    dv_num = DataValidation(
        type="whole", operator="greaterThanOrEqual", formula1="0",
        allow_blank=True, showInputMessage=True, showErrorMessage=True,
        promptTitle="数量入力", prompt="設備の個数を入力（0以上の整数）",
        errorTitle="入力エラー", error="0以上の整数を入力してください"
    )
    has_num = False
    for row in range(5, 28):
        for col in range(3, 15):  # C=3 to N=14
            cl = get_column_letter(col)
            cell = ws[f"{cl}{row}"]
            val = cell.value
            if isinstance(val, str) and val.startswith('='):
                continue
            dv_num.add(cell)
            cell.protection = unlocked
            has_num = True
    if has_num:
        ws.add_data_validation(dv_num)
    
    # Unlock name/number entry cells (A, B columns)
    for row in range(5, 28):
        for cl in ['A', 'B']:
            cell = ws[f"{cl}{row}"]
            val = cell.value
            if isinstance(val, str) and val.startswith('='):
                continue
            cell.protection = unlocked
    
    # Unlock remarks cells (Q column)
    for row in range(5, 28):
        ws[f"Q{row}"].protection = unlocked
    
    # Protect the sheet
    ws.protection = SheetProtection(
        sheet=True, objects=True, scenarios=True,
        formatCells=False, formatColumns=False, formatRows=False,
        selectLockedCells=False, selectUnlockedCells=False
    )


files = [
    ("/root/.claude/uploads/50898019-d810-5420-b368-3e2dd0abe1a5/bcbf0df6-___________.xlsx", "消防設備試験結果"),
    ("/root/.claude/uploads/50898019-d810-5420-b368-3e2dd0abe1a5/9145e47d-___________________.xlsx", "共同住宅用自動火災報知設備"),
    ("/root/.claude/uploads/50898019-d810-5420-b368-3e2dd0abe1a5/0abf8a6e-______.xlsx", "試験器"),
]

skip_sheets = {'Sheet1', 'Sheet2'}

dst_wb = openpyxl.Workbook()
dst_wb.remove(dst_wb.active)

used_names = set()

for fpath, file_label in files:
    src_wb = openpyxl.load_workbook(fpath)
    for sn in src_wb.sheetnames:
        if sn in skip_sheets:
            continue
        
        if file_label == "共同住宅用自動火災報知設備":
            new_name = f"火報_{sn}"
        elif file_label == "試験器":
            new_name = "試験器"
        else:
            new_name = sn
        
        new_name = new_name[:31]
        if new_name in used_names:
            new_name = new_name[:28] + "_2"
        used_names.add(new_name)
        
        print(f"Copying: {file_label} / {sn} -> {new_name}")
        ws = copy_sheet(src_wb[sn], dst_wb, new_name)
        
        # Apply input improvements to fire alarm detail sheets
        if file_label == "共同住宅用自動火災報知設備" and sn not in skip_sheets:
            add_input_improvements(ws, "fire_alarm_detail")

output_path = "/home/user/claude-project/消防設備_統合ファイル.xlsx"
dst_wb.save(output_path)
print(f"\nSaved: {output_path}")
print(f"Sheets ({len(dst_wb.sheetnames)}): {dst_wb.sheetnames}")
