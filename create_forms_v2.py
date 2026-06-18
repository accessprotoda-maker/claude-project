#!/usr/bin/env python3
"""
消防設備工事用書類 Excel生成スクリプト v2
着工届出書はPDFレイアウトに合わせて手動構築、他シートはWord自動変換（改良版）
"""
import sys
from docx import Document
from docx.oxml.ns import qn
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ============================================================
# Styles
# ============================================================
F_TITLE = Font(name='ＭＳ ゴシック', bold=True, size=14)
F_HEADER = Font(name='ＭＳ ゴシック', bold=True, size=9)
F_NORMAL = Font(name='ＭＳ ゴシック', size=9)
F_SMALL = Font(name='ＭＳ ゴシック', size=7.5)
F_SECTION = Font(name='ＭＳ ゴシック', bold=True, size=11)

THIN = Border(
    left=Side('thin'), right=Side('thin'),
    top=Side('thin'), bottom=Side('thin'))
NO_BORDER = Border()

FILL_INPUT = PatternFill('solid', fgColor='FFFFCC')
FILL_AUTO = PatternFill('solid', fgColor='E0FFE0')
FILL_HEADER = PatternFill('solid', fgColor='D9E1F2')
FILL_SECTION = PatternFill('solid', fgColor='B4C6E7')

AL_C = Alignment(horizontal='center', vertical='center', wrap_text=True)
AL_L = Alignment(horizontal='left', vertical='center', wrap_text=True)
AL_R = Alignment(horizontal='right', vertical='center', wrap_text=True)
AL_TL = Alignment(horizontal='left', vertical='top', wrap_text=True)
AL_CV = Alignment(horizontal='center', vertical='center', wrap_text=True, text_rotation=255)

def sc(ws, r, c, val=None, font=F_NORMAL, al=AL_C, border=THIN, fill=None):
    cell = ws.cell(row=r, column=c, value=val)
    if font: cell.font = font
    if al: cell.alignment = al
    if border: cell.border = border
    if fill: cell.fill = fill
    return cell

def input_cell(ws, r, c, val=None):
    return sc(ws, r, c, val, fill=FILL_INPUT)

def header_cell(ws, r, c, val=None):
    return sc(ws, r, c, val, font=F_HEADER, fill=FILL_HEADER)

def merge_and_set(ws, r1, c1, r2, c2, val=None, font=F_NORMAL, al=AL_C, border=THIN, fill=None):
    if r1 != r2 or c1 != c2:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    sc(ws, r1, c1, val, font, al, border, fill)
    for r in range(r1, r2+1):
        for c in range(c1, c2+1):
            ws.cell(row=r, column=c).border = border

def merge_no_border(ws, r1, c1, r2, c2, val=None, font=F_NORMAL, al=AL_C):
    if r1 != r2 or c1 != c2:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = ws.cell(row=r1, column=c1, value=val)
    if font: cell.font = font
    if al: cell.alignment = al
    return cell


# ============================================================
# Word table -> Excel conversion engine (improved)
# ============================================================

def get_grid_col_widths(table):
    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is None:
        return []
    cols = grid.findall(qn('w:gridCol'))
    return [int(c.get(qn('w:w'), '0')) for c in cols]

def emu_to_excel_width(emu_widths, target_total=85, min_width=3):
    """Convert EMU/twip widths to Excel character widths with minimum width."""
    total = sum(emu_widths)
    if total == 0:
        return [10] * len(emu_widths)
    result = [(w / total) * target_total for w in emu_widths]
    # Enforce minimum width
    for i in range(len(result)):
        if result[i] < min_width:
            result[i] = min_width
    return result

def parse_table_cells(table):
    rows = []
    for tr in table._tbl.findall(qn('w:tr')):
        row_cells = []
        col_idx = 0
        for tc in tr.findall(qn('w:tc')):
            tcPr = tc.find(qn('w:tcPr'))
            gridSpan_elem = tcPr.find(qn('w:gridSpan')) if tcPr is not None else None
            span = int(gridSpan_elem.get(qn('w:val'))) if gridSpan_elem is not None else 1
            vMerge_elem = tcPr.find(qn('w:vMerge')) if tcPr is not None else None
            is_vmerge_start = False
            is_vmerge_continue = False
            if vMerge_elem is not None:
                val = vMerge_elem.get(qn('w:val'))
                if val == 'restart':
                    is_vmerge_start = True
                else:
                    is_vmerge_continue = True
            paragraphs = tc.findall(qn('w:p'))
            texts = []
            for p in paragraphs:
                p_text = ''
                for node in p.iter():
                    if node.tag == qn('w:t'):
                        p_text += (node.text or '')
                texts.append(p_text)
            text = '\n'.join(texts).strip()
            row_cells.append({
                'col_idx': col_idx, 'span': span, 'text': text,
                'is_vmerge_start': is_vmerge_start, 'is_vmerge_continue': is_vmerge_continue,
            })
            col_idx += span
        rows.append(row_cells)
    return rows

def compute_vmerge_ranges(parsed_rows):
    active = {}
    vmerge_map = {}
    for row_idx, row in enumerate(parsed_rows):
        for cell in row:
            ci = cell['col_idx']
            if cell['is_vmerge_start']:
                active[ci] = row_idx
                vmerge_map[(row_idx, ci)] = row_idx
            elif cell['is_vmerge_continue']:
                if ci in active:
                    vmerge_map[(active[ci], ci)] = row_idx
            else:
                if ci in active:
                    del active[ci]
    return vmerge_map

def write_table_to_sheet(ws, table, start_row=1, target_width=85):
    grid_widths = get_grid_col_widths(table)
    num_grid_cols = len(grid_widths)
    excel_widths = emu_to_excel_width(grid_widths, target_width, min_width=3)
    for i, w in enumerate(excel_widths):
        col_letter = get_column_letter(i + 1)
        if ws.column_dimensions[col_letter].width is None or ws.column_dimensions[col_letter].width < w:
            ws.column_dimensions[col_letter].width = w
    parsed_rows = parse_table_cells(table)
    vmerge_map = compute_vmerge_ranges(parsed_rows)

    # Check if first row is a full-width merged cell with multi-line content
    header_rows_added = 0
    if parsed_rows and len(parsed_rows[0]) == 1:
        cell0 = parsed_rows[0][0]
        if cell0['span'] >= num_grid_cols and '\n' in cell0['text']:
            # Split into separate rows above the table
            lines = cell0['text'].split('\n')
            for line in lines:
                line = line.strip()
                if line:
                    c = ws.cell(row=start_row + header_rows_added, column=1, value=line)
                    c.font = F_NORMAL
                    c.alignment = AL_L
                    if num_grid_cols > 1:
                        ws.merge_cells(start_row=start_row + header_rows_added, start_column=1,
                                       end_row=start_row + header_rows_added, end_column=num_grid_cols)
                    header_rows_added += 1
            header_rows_added += 1  # blank row
            parsed_rows = parsed_rows[1:]  # skip the first row
            vmerge_map = compute_vmerge_ranges(parsed_rows)

    actual_start = start_row + header_rows_added

    for row_idx, row in enumerate(parsed_rows):
        excel_row = actual_start + row_idx
        for cell in row:
            ci = cell['col_idx']
            excel_col = ci + 1
            span = cell['span']
            if cell['is_vmerge_continue']:
                for c in range(excel_col, excel_col + span):
                    ws.cell(row=excel_row, column=c).border = THIN
                    ws.cell(row=excel_row, column=c).font = F_NORMAL
                continue
            vmerge_end_row = row_idx
            if cell['is_vmerge_start'] and (row_idx, ci) in vmerge_map:
                vmerge_end_row = vmerge_map[(row_idx, ci)]
            end_excel_row = actual_start + vmerge_end_row
            end_excel_col = excel_col + span - 1
            if end_excel_row > excel_row or end_excel_col > excel_col:
                try:
                    ws.merge_cells(start_row=excel_row, start_column=excel_col,
                                   end_row=end_excel_row, end_column=end_excel_col)
                except Exception:
                    pass
            text = cell['text']
            c = ws.cell(row=excel_row, column=excel_col)
            c.value = text if text else None
            c.font = F_NORMAL
            c.alignment = AL_C
            c.border = THIN
            for rr in range(excel_row, end_excel_row + 1):
                for cc in range(excel_col, end_excel_col + 1):
                    ws.cell(row=rr, column=cc).border = THIN
                    ws.cell(row=rr, column=cc).font = F_NORMAL

        max_lines = 1
        for cell in row:
            if not cell['is_vmerge_continue']:
                lines = cell['text'].count('\n') + 1 if cell['text'] else 1
                max_lines = max(max_lines, lines)
        ws.row_dimensions[excel_row].height = max(15, min(max_lines * 13, 60))

    return actual_start + len(parsed_rows)

def setup_print_area(ws, max_row, max_col):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.print_area = f'A1:{get_column_letter(max_col)}{max_row}'


# ============================================================
# Source files
# ============================================================
FILES = {
    '設置届': '/root/.claude/uploads/1aedaef9-61f4-5511-9d57-68b254cd1569/a3c02a33-_____.docx',
    '概要表': '/tmp/docconv/dc729bb0-_____.docx',
    '様式11': '/tmp/docconv/18d9644f-sikenkekkahoukoku11.docx',
    '様式34': '/tmp/docconv/5b588e9e-__sikenkekkahoukoku34.docx',
}

OUTPUT = '/home/user/claude-project/消防設備工事_様式統合.xlsx'


def build_input_sheet(wb):
    """入力シート - copied from convert_word_to_excel.py"""
    ws = wb.active
    ws.title = '入力シート'
    ws.sheet_properties.tabColor = 'FF0000'

    for c in range(1, 8):
        ws.column_dimensions[get_column_letter(c)].width = 28

    merge_and_set(ws, 1, 1, 1, 7, '消防設備工事 統合入力シート', F_TITLE, AL_C)
    sc(ws, 2, 1, '※ 黄色セルに入力 → 各書式に自動反映されます', F_SMALL, AL_L)

    r = 4
    merge_and_set(ws, r, 1, r, 4, '【基本情報】', F_SECTION, AL_L, fill=FILL_SECTION)

    fields = [
        ('担当消防署（例：○○消防署長　殿）', None),
        ('届出日', None),
        ('届出者 住所', None),
        ('届出者 氏名（管理組合名等）', None),
        ('工事場所（住所）', None),
        ('防火対象物の名称', None),
        ('設備の種類', '自動火災報知設備'),
        ('用途（例：5項ロ）', None),
        ('構造', '耐火構造'),
        ('地上階数', None),
        ('地下階数', None),
        ('延べ面積（㎡）', None),
        ('着工予定日', None),
        ('完成予定日', None),
        ('理事長名（設置届用）', None),
        ('全住戸数', None),
    ]
    for i, (label, default) in enumerate(fields):
        rr = 5 + i
        header_cell(ws, rr, 1, label)
        input_cell(ws, rr, 2, default)

    dv = DataValidation(type="list", formula1='"自動火災報知設備,共同住宅用自動火災報知設備,住戸用自動火災報知設備"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(ws['B11'])

    dv2 = DataValidation(type="list", formula1='"耐火構造,準耐火構造,その他"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(ws['B13'])

    r = 22
    merge_and_set(ws, r, 1, r, 4, '【施工者情報】', F_SECTION, AL_L, fill=FILL_SECTION)
    for i, (label, default) in enumerate([
        ('施工会社名', '株式会社　TSCアクセス・プロ'),
        ('代表者名', '代表取締役　中村　勇'),
        ('施工会社住所', '名古屋市中区栄４丁目１６番３号　山光堂ビル５階'),
        ('電話番号', '052-269-9100'),
    ]):
        header_cell(ws, 23+i, 1, label)
        input_cell(ws, 23+i, 2, default)

    r = 28
    merge_and_set(ws, r, 1, r, 7, '【消防設備士マスタ】', F_SECTION, AL_L, fill=FILL_SECTION)
    mst_headers = ['氏名','住所','免状種類','交付年月日','交付番号','講習年月','交付知事']
    for i, h in enumerate(mst_headers):
        header_cell(ws, 29, 1+i, h)

    staff = [
        ['河合　佑樹','愛知県尾張旭市晴丘町東236番地4','甲4','R5年7月21日','第00324号','','愛知'],
        ['戸田　健仁','愛知県長久手市西原山1-1 ｾﾝﾄｱｰｽ502','甲4','R5年11月7日','第00448号','','愛知'],
        ['棚町　征','岐阜県土岐市泉町大富195番地の249','甲4','H14年10月30日','第00009号','令和元年10月','愛知'],
    ]
    for i, s in enumerate(staff):
        for j, v in enumerate(s):
            sc(ws, 30+i, 1+j, v, border=THIN)

    r = 34
    header_cell(ws, r, 1, '着工届 届出者（選択）')
    input_cell(ws, r, 2, '戸田　健仁')
    header_cell(ws, r+1, 1, '試験実施者（選択）')
    input_cell(ws, r+1, 2)
    header_cell(ws, r+2, 1, '配線試験実施者（選択）')
    input_cell(ws, r+2, 2, '戸田　健仁')

    for rr in range(34, 37):
        dv_s = DataValidation(type="list", formula1='"河合　佑樹,戸田　健仁,棚町　征"', allow_blank=True)
        ws.add_data_validation(dv_s)
        dv_s.add(ws.cell(row=rr, column=2))

    r = 38
    merge_and_set(ws, r, 1, r, 4, '【受信機・機器情報（共用部）】', F_SECTION, AL_L, fill=FILL_SECTION)
    for i, (label, default) in enumerate([
        ('受信機 型級（P・GP型 ○級）', '1'),
        ('回線数（例：10/10）', '10/10'),
        ('予備電源 DC（V）', '24'),
        ('予備電源（AH）', '0.45'),
        ('受信機設置場所', '管理人室'),
        ('受信機メーカー', 'パナソニック㈱'),
        ('型式番号（受第○～○号）', ''),
    ]):
        header_cell(ws, 39+i, 1, label)
        input_cell(ws, 39+i, 2, default)

    r = 47
    merge_and_set(ws, r, 1, r, 4, '【受信機・機器情報（住戸部）】', F_SECTION, AL_L, fill=FILL_SECTION)
    for i, (label, default) in enumerate([
        ('受信機 型級（P・GP型 ○級）', '3'),
        ('回線数（例：1/1）', '1/1'),
        ('受信機設置場所', '各住戸リビング'),
        ('受信機メーカー', 'パナソニック㈱'),
        ('型式番号（受第○～○号）', ''),
        ('表示器台数', ''),
    ]):
        header_cell(ws, 48+i, 1, label)
        input_cell(ws, 48+i, 2, default)

    r = 55
    merge_and_set(ws, r, 1, r, 4, '【地区音響装置・発信機・表示灯】', F_SECTION, AL_L, fill=FILL_SECTION)
    for i, (label, default) in enumerate([
        ('地区音響 認評音第（号）', '13-5'),
        ('地区音響 メーカー', 'パナソニック㈱'),
        ('地区音響 電圧DC（V）', '24'),
        ('地区音響 電流（mA）', '10'),
        ('地区音響 個数', ''),
        ('地区音響 音量（dB）', '95'),
        ('地区音響 鐘径（mm）', '150'),
        ('発信機 型（P等）', 'P'),
        ('発信機 級', '1'),
        ('発信機 個数（屋外型）', ''),
        ('発信機 型式番号（発第）', ''),
        ('発信機 メーカー', 'パナソニック㈱'),
        ('表示灯 電圧DC（V）', '24'),
        ('表示灯 電流（mA）', '9'),
        ('表示灯 個数', ''),
    ]):
        header_cell(ws, 56+i, 1, label)
        input_cell(ws, 56+i, 2, default)

    r = 72
    merge_and_set(ws, r, 1, r, 4, '【配線試験 絶縁抵抗値】', F_SECTION, AL_L, fill=FILL_SECTION)
    for i, (label, default) in enumerate([
        ('電源回路 試験電圧（V）', '250'),
        ('電源回路 絶縁抵抗（MΩ）', '150'),
        ('操作回路 試験電圧（V）', '50'),
        ('操作回路 絶縁抵抗（MΩ）', '20'),
        ('表示灯回路 試験電圧（V）', '50'),
        ('表示灯回路 絶縁抵抗（MΩ）', '20'),
        ('警報回路 試験電圧（V）', '50'),
        ('警報回路 絶縁抵抗（MΩ）', '20'),
        ('感知器回路 試験電圧（V）', '50'),
        ('感知器回路 絶縁抵抗（MΩ）', '20'),
    ]):
        header_cell(ws, 73+i, 1, label)
        input_cell(ws, 73+i, 2, default)

    r = 84
    merge_and_set(ws, r, 1, r, 4, '【工事関連】', F_SECTION, AL_L, fill=FILL_SECTION)
    header_cell(ws, 85, 1, '工事種別')
    input_cell(ws, 85, 2, '取替え')
    dv_w = DataValidation(type="list", formula1='"新設,増設,移設,取替え,改造,その他"', allow_blank=True)
    ws.add_data_validation(dv_w)
    dv_w.add(ws['B85'])

    header_cell(ws, 86, 1, '試験実施日')
    input_cell(ws, 86, 2)
    header_cell(ws, 87, 1, '完成年月日')
    input_cell(ws, 87, 2)

    header_cell(ws, 88, 1, '関連設備')
    input_cell(ws, 88, 2, 'インターホン設備')

    header_cell(ws, 89, 1, '鳴動方式')
    input_cell(ws, 89, 2, '一斉鳴動')
    dv_ring = DataValidation(type="list", formula1='"一斉鳴動,区分鳴動"', allow_blank=True)
    ws.add_data_validation(dv_ring)
    dv_ring.add(ws['B89'])

    header_cell(ws, 90, 1, '放送設備連動')
    input_cell(ws, 90, 2, '無')
    dv_yn = DataValidation(type="list", formula1='"有,無"', allow_blank=True)
    ws.add_data_validation(dv_yn)
    dv_yn.add(ws['B90'])

    header_cell(ws, 91, 1, '中継器設置場所')
    input_cell(ws, 91, 2, '戸外表示機組込')

    header_cell(ws, 92, 1, '受信機操作部 高さ（m）')
    input_cell(ws, 92, 2, '1.4')

    setup_print_area(ws, 92, 7)
    return ws


def build_chakko_todoke(wb):
    """着工届出書 - manually built to match PDF layout exactly."""
    ws = wb.create_sheet('着工届出書')

    # Column widths: A-J (10 columns)
    # A=4, B=14, C=3, D=7, E=9, F=6, G=9, H=6, I=7, J=9
    widths = [4, 14, 3, 7, 9, 6, 9, 6, 7, 9]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i+1)].width = w

    # ========== PRE-TABLE (no borders) ==========
    # Row 1: 別記様式
    merge_no_border(ws, 1, 1, 1, 10, '別記様式第１号の７（第33条の18関係）', F_SMALL, AL_L)
    ws.row_dimensions[1].height = 15

    # Row 2: empty
    ws.row_dimensions[2].height = 8

    # Row 3: Title
    merge_no_border(ws, 3, 1, 3, 10, '工事整備対象設備等着工届出書', F_TITLE, AL_C)
    ws.row_dimensions[3].height = 25

    # Row 4: empty
    ws.row_dimensions[4].height = 8

    # Row 5: Date (right-aligned) - formula linked to 入力シート
    merge_no_border(ws, 5, 7, 5, 10, None, F_NORMAL, AL_R)
    ws.cell(row=5, column=7).value = '=入力シート!B6'
    ws.cell(row=5, column=7).font = F_NORMAL
    ws.cell(row=5, column=7).alignment = AL_R
    ws.row_dimensions[5].height = 18

    # Row 6: empty
    ws.row_dimensions[6].height = 8

    # Row 7: 消防署長 殿
    merge_no_border(ws, 7, 1, 7, 5, None, F_NORMAL, AL_L)
    ws.cell(row=7, column=1).value = '=入力シート!B5&"　殿"'
    ws.cell(row=7, column=1).font = F_NORMAL
    ws.cell(row=7, column=1).alignment = AL_L
    ws.row_dimensions[7].height = 18

    # Row 8: empty
    ws.row_dimensions[8].height = 8

    # Row 9: 届出者 label (center-right)
    merge_no_border(ws, 9, 6, 9, 8, '届　出　者', F_NORMAL, AL_C)
    ws.row_dimensions[9].height = 18

    # Row 10: 住所
    merge_no_border(ws, 10, 5, 10, 5, '住　所', F_NORMAL, AL_R)
    merge_no_border(ws, 10, 6, 10, 10, None, F_NORMAL, AL_L)
    ws.cell(row=10, column=6).value = '=VLOOKUP(入力シート!B34,入力シート!A30:G32,2,FALSE)'
    ws.row_dimensions[10].height = 18

    # Row 11: 氏名
    merge_no_border(ws, 11, 5, 11, 5, '氏　名', F_NORMAL, AL_R)
    merge_no_border(ws, 11, 6, 11, 10, None, F_NORMAL, AL_L)
    ws.cell(row=11, column=6).value = '=入力シート!B34'
    ws.row_dimensions[11].height = 18

    # Row 12: empty
    ws.row_dimensions[12].height = 8

    # ========== TABLE AREA (with borders) ==========
    TR = 13  # table start row

    # Row 13: 工事の場所
    merge_and_set(ws, TR, 1, TR, 2, '工　事　の　場　所', F_NORMAL, AL_C)
    merge_and_set(ws, TR, 3, TR, 10, None, F_NORMAL, AL_L)
    ws.cell(row=TR, column=3).value = '=入力シート!B9'
    ws.row_dimensions[TR].height = 20

    # Row 14: 防火対象物の名称
    merge_and_set(ws, TR+1, 1, TR+1, 2, '工事を行う防火対象物の名称', F_NORMAL, AL_C)
    merge_and_set(ws, TR+1, 3, TR+1, 10, None, F_NORMAL, AL_L)
    ws.cell(row=TR+1, column=3).value = '=入力シート!B10'
    ws.row_dimensions[TR+1].height = 20

    # Row 15: 設備の種類
    merge_and_set(ws, TR+2, 1, TR+2, 2, '工事整備対象設備等の種類', F_NORMAL, AL_C)
    merge_and_set(ws, TR+2, 3, TR+2, 10, None, F_NORMAL, AL_L)
    ws.cell(row=TR+2, column=3).value = '=入力シート!B11'
    ws.row_dimensions[TR+2].height = 20

    # Rows 16-17: 施工者 (A spans 2 rows vertically)
    merge_and_set(ws, TR+3, 1, TR+4, 1, '工事整備対象設備等の工事施工者', F_SMALL, AL_CV)
    # Row 16: 住所
    merge_and_set(ws, TR+3, 2, TR+3, 2, '住　所', F_NORMAL, AL_C)
    merge_and_set(ws, TR+3, 3, TR+3, 10, None, F_NORMAL, AL_L)
    ws.cell(row=TR+3, column=3).value = '=入力シート!B25&"　電話番号　"&入力シート!B26'
    ws.row_dimensions[TR+3].height = 22

    # Row 17: 氏名
    merge_and_set(ws, TR+4, 2, TR+4, 2, '氏　名\n法人の場合は名称\n及び代表者氏名', F_SMALL, AL_C)
    merge_and_set(ws, TR+4, 3, TR+4, 10, None, F_NORMAL, AL_L)
    ws.cell(row=TR+4, column=3).value = '=入力シート!B23&"　"&入力シート!B24'
    ws.row_dimensions[TR+4].height = 40

    # Rows 18-21: 消防設備士 section
    merge_and_set(ws, TR+5, 1, TR+8, 1, '消防設備士', F_NORMAL, AL_CV)

    # Row 18: header line 1
    merge_and_set(ws, TR+5, 2, TR+5, 3, '免　状　の', F_NORMAL, AL_C)
    merge_and_set(ws, TR+5, 4, TR+5, 4, '種類等', F_SMALL, AL_C)
    merge_and_set(ws, TR+5, 5, TR+5, 5, '交付知事', F_NORMAL, AL_C)
    merge_and_set(ws, TR+5, 6, TR+5, 7, '交付年月日', F_NORMAL, AL_C)
    merge_and_set(ws, TR+5, 8, TR+5, 10, '講習受講状況', F_NORMAL, AL_C)
    ws.row_dimensions[TR+5].height = 18

    # Row 19: header line 2
    merge_and_set(ws, TR+6, 2, TR+6, 3, '種類及び\n指定区分', F_SMALL, AL_C)
    merge_and_set(ws, TR+6, 4, TR+6, 4, '', F_NORMAL, AL_C)  # blank (merged concept w/ above conceptually)
    merge_and_set(ws, TR+6, 5, TR+6, 5, '', F_NORMAL, AL_C)
    merge_and_set(ws, TR+6, 6, TR+6, 7, '交付番号', F_NORMAL, AL_C)
    merge_and_set(ws, TR+6, 8, TR+6, 8, '受講地', F_NORMAL, AL_C)
    merge_and_set(ws, TR+6, 9, TR+6, 10, '講習年月', F_NORMAL, AL_C)
    ws.row_dimensions[TR+6].height = 25

    # Row 20: data line 1 (甲)
    merge_and_set(ws, TR+7, 2, TR+7, 2, '甲', F_NORMAL, AL_C)
    # C20: 種4類 from VLOOKUP
    merge_and_set(ws, TR+7, 3, TR+7, 3, None, F_NORMAL, AL_C)
    ws.cell(row=TR+7, column=3).value = '=MID(VLOOKUP(入力シート!B34,入力シート!A30:G32,3,FALSE),2,10)'
    merge_and_set(ws, TR+7, 4, TR+7, 4, '', F_NORMAL, AL_C)
    # E20: 交付知事
    merge_and_set(ws, TR+7, 5, TR+7, 5, None, F_NORMAL, AL_C)
    ws.cell(row=TR+7, column=5).value = '=VLOOKUP(入力シート!B34,入力シート!A30:G32,7,FALSE)'
    # F-G: 交付年月日
    merge_and_set(ws, TR+7, 6, TR+7, 7, None, F_NORMAL, AL_C)
    ws.cell(row=TR+7, column=6).value = '=VLOOKUP(入力シート!B34,入力シート!A30:G32,4,FALSE)'
    # H: 講習受講地
    merge_and_set(ws, TR+7, 8, TR+7, 8, None, F_NORMAL, AL_C)
    ws.cell(row=TR+7, column=8).value = '=VLOOKUP(入力シート!B34,入力シート!A30:G32,7,FALSE)'
    # I-J: 講習年月
    merge_and_set(ws, TR+7, 9, TR+7, 10, None, F_NORMAL, AL_C)
    ws.cell(row=TR+7, column=9).value = '=VLOOKUP(入力シート!B34,入力シート!A30:G32,6,FALSE)'
    ws.row_dimensions[TR+7].height = 20

    # Row 21: data line 2 (・乙)
    merge_and_set(ws, TR+8, 2, TR+8, 2, '・\n乙', F_NORMAL, AL_C)
    merge_and_set(ws, TR+8, 3, TR+8, 3, '', F_NORMAL, AL_C)
    merge_and_set(ws, TR+8, 4, TR+8, 4, '', F_NORMAL, AL_C)
    merge_and_set(ws, TR+8, 5, TR+8, 5, '', F_NORMAL, AL_C)
    # F-G: 交付番号
    merge_and_set(ws, TR+8, 6, TR+8, 7, None, F_NORMAL, AL_C)
    ws.cell(row=TR+8, column=6).value = '=VLOOKUP(入力シート!B34,入力シート!A30:G32,5,FALSE)'
    merge_and_set(ws, TR+8, 8, TR+8, 8, '', F_NORMAL, AL_C)
    merge_and_set(ws, TR+8, 9, TR+8, 10, '', F_NORMAL, AL_C)
    ws.row_dimensions[TR+8].height = 25

    # Rows 22-23: 工事の種別
    merge_and_set(ws, TR+9, 1, TR+10, 4, '工　事　の　種　別', F_NORMAL, AL_C)
    merge_and_set(ws, TR+9, 5, TR+9, 10, '１　新設　　２　増設　　３　移設　　４　取替え', F_NORMAL, AL_L)
    merge_and_set(ws, TR+10, 5, TR+10, 10, '５　改造　　６　その他', F_NORMAL, AL_L)
    ws.row_dimensions[TR+9].height = 20
    ws.row_dimensions[TR+10].height = 20

    # Row 24: 着工予定日 / 完成予定日
    merge_and_set(ws, TR+11, 1, TR+11, 4, '着　工　予　定　日', F_NORMAL, AL_C)
    merge_and_set(ws, TR+11, 5, TR+11, 6, None, F_NORMAL, AL_C)
    ws.cell(row=TR+11, column=5).value = '=入力シート!B17'
    merge_and_set(ws, TR+11, 7, TR+11, 8, '完　成　予　定　日', F_NORMAL, AL_C)
    merge_and_set(ws, TR+11, 9, TR+11, 10, None, F_NORMAL, AL_C)
    ws.cell(row=TR+11, column=9).value = '=入力シート!B18'
    ws.row_dimensions[TR+11].height = 20

    # Rows 25-28: 受付欄 / 経過欄
    merge_and_set(ws, TR+12, 1, TR+13, 6, '※　受　付　欄', F_NORMAL, AL_C)
    merge_and_set(ws, TR+12, 7, TR+13, 10, '※　経　過　欄', F_NORMAL, AL_C)
    ws.row_dimensions[TR+12].height = 18

    merge_and_set(ws, TR+14, 1, TR+15, 6, '', F_NORMAL, AL_C)
    merge_and_set(ws, TR+14, 7, TR+15, 10, '', F_NORMAL, AL_C)
    ws.row_dimensions[TR+13].height = 30
    ws.row_dimensions[TR+14].height = 30
    ws.row_dimensions[TR+15].height = 30

    # ========== POST-TABLE (no borders) ==========
    PR = TR + 17  # post row start (row 30)
    ws.row_dimensions[TR+16].height = 8  # blank row

    merge_no_border(ws, PR, 1, PR, 10,
                    '備考　１　この用紙の大きさは、日本産業規格A４とすること。', F_SMALL, AL_L)
    merge_no_border(ws, PR+1, 1, PR+1, 10,
                    '　　　２　工事の種別の欄は、該当する事項を○印で囲むこと。', F_SMALL, AL_L)
    merge_no_border(ws, PR+2, 1, PR+2, 10,
                    '　　　３　※印の欄には、記入しないこと。', F_SMALL, AL_L)

    setup_print_area(ws, PR+2, 10)
    return ws


def build_setchi_todoke(wb):
    """設置届出書 - auto-converted from Word with improved widths."""
    print("Processing 設置届出書...")
    doc = Document(FILES['設置届'])
    ws = wb.create_sheet('設置届出書')

    pre_paras = []
    for elem in doc.element.body:
        if elem.tag == qn('w:tbl'):
            break
        if elem.tag == qn('w:p'):
            text = ''.join(t.text or '' for t in elem.iter() if t.tag == qn('w:t'))
            if text.strip():
                pre_paras.append(text.strip())

    row = 1
    for text in pre_paras:
        c = ws.cell(row=row, column=1, value=text)
        c.font = F_NORMAL
        c.alignment = AL_L
        row += 1
    if pre_paras:
        row += 1

    next_row = write_table_to_sheet(ws, doc.tables[0], start_row=row)

    post_paras = []
    found_table = False
    for elem in doc.element.body:
        if elem.tag == qn('w:tbl'):
            found_table = True
            continue
        if found_table and elem.tag == qn('w:p'):
            text = ''.join(t.text or '' for t in elem.iter() if t.tag == qn('w:t'))
            if text.strip():
                post_paras.append(text.strip())

    next_row += 1
    for text in post_paras:
        c = ws.cell(row=next_row, column=1, value=text)
        c.font = F_SMALL
        c.alignment = AL_L
        next_row += 1

    grid_widths = get_grid_col_widths(doc.tables[0])
    setup_print_area(ws, next_row, len(grid_widths))
    return ws


def build_gaiyo_sheets(wb):
    """概要表 - 4 sheets: その1共用部, その2共用部, その1住戸部, その2住戸部"""
    print("Processing 概要表...")
    doc = Document(FILES['概要表'])

    body_elements = list(doc.element.body)
    table_indices = [i for i, elem in enumerate(body_elements) if elem.tag == qn('w:tbl')]

    sheet_configs = [
        ('概要表（その1）共用部', 0),
        ('概要表（その2）共用部', 1),
        ('概要表（その1）住戸部', 0),
        ('概要表（その2）住戸部', 1),
    ]

    for sheet_name, tbl_num in sheet_configs:
        ws_g = wb.create_sheet(sheet_name)

        if tbl_num < len(table_indices):
            tbl_idx = table_indices[tbl_num]
            start_search = table_indices[tbl_num - 1] + 1 if tbl_num > 0 else 0
            pre_texts = []
            for i in range(start_search, tbl_idx):
                elem = body_elements[i]
                if elem.tag == qn('w:p'):
                    text = ''.join(t.text or '' for t in elem.iter() if t.tag == qn('w:t'))
                    if text.strip():
                        pre_texts.append(text.strip())

            row = 1
            # Add 共用部/住戸部 label
            suffix = '共用部' if '共用部' in sheet_name else '住戸部'
            c = ws_g.cell(row=row, column=1, value=f'【{suffix}】')
            c.font = F_SECTION
            c.alignment = AL_L
            row += 1

            for text in pre_texts:
                c = ws_g.cell(row=row, column=1, value=text)
                c.font = F_NORMAL
                c.alignment = AL_L
                row += 1
            if pre_texts:
                row += 1

            next_row = write_table_to_sheet(ws_g, doc.tables[tbl_num], start_row=row)
            grid_widths = get_grid_col_widths(doc.tables[tbl_num])
            setup_print_area(ws_g, next_row, len(grid_widths))

    # Post-table paragraphs on last sheet
    if len(table_indices) >= 2:
        last_tbl_idx = table_indices[-1]
        post_texts = []
        for i in range(last_tbl_idx + 1, len(body_elements)):
            elem = body_elements[i]
            if elem.tag == qn('w:p'):
                text = ''.join(t.text or '' for t in elem.iter() if t.tag == qn('w:t'))
                if text.strip():
                    post_texts.append(text.strip())
        if post_texts:
            # Add to the last created sheet
            ws_last = wb['概要表（その2）住戸部']
            # Find last used row
            last_row = ws_last.max_row + 1
            for text in post_texts:
                c = ws_last.cell(row=last_row, column=1, value=text)
                c.font = F_SMALL
                c.alignment = AL_L
                last_row += 1


def build_shiken_sheet(wb, file_key, sheet_name):
    """試験結果 sheets - auto-converted from Word."""
    print(f"Processing {sheet_name}...")
    doc = Document(FILES[file_key])
    ws = wb.create_sheet(sheet_name)

    body_elements = list(doc.element.body)
    row = 1
    current_table_idx = 0

    for elem in body_elements:
        if elem.tag == qn('w:p'):
            text = ''.join(t.text or '' for t in elem.iter() if t.tag == qn('w:t'))
            if text.strip():
                c = ws.cell(row=row, column=1, value=text.strip())
                c.font = F_NORMAL
                c.alignment = AL_L
                row += 1
        elif elem.tag == qn('w:tbl') and current_table_idx < len(doc.tables):
            table = doc.tables[current_table_idx]
            grid_widths = get_grid_col_widths(table)
            excel_widths = emu_to_excel_width(grid_widths, 85, min_width=3)
            for i, w in enumerate(excel_widths):
                col_letter = get_column_letter(i + 1)
                existing = ws.column_dimensions[col_letter].width
                if existing is None or existing < w:
                    ws.column_dimensions[col_letter].width = w

            row = write_table_to_sheet(ws, table, start_row=row, target_width=85)
            row += 1
            current_table_idx += 1

    if doc.tables:
        max_cols = max(len(get_grid_col_widths(t)) for t in doc.tables)
        setup_print_area(ws, row, max_cols)


def main():
    wb = openpyxl.Workbook()

    # Sheet 1: 入力シート
    build_input_sheet(wb)

    # Sheet 2: 着工届出書 (manually built to match PDF)
    build_chakko_todoke(wb)

    # Sheet 3: 設置届出書
    build_setchi_todoke(wb)

    # Sheets 4-7: 概要表 (4 sheets)
    build_gaiyo_sheets(wb)

    # Sheet 8: 試験結果(様式11)
    build_shiken_sheet(wb, '様式11', '試験結果(様式11)')

    # Sheet 9: 試験結果(様式34)
    build_shiken_sheet(wb, '様式34', '試験結果(様式34)')

    # Save
    wb.save(OUTPUT)
    print(f"Saved to {OUTPUT}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == '__main__':
    main()
